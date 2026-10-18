"""Extract published tracrRNAs without treating predictions as experimental truth.

Run with the artifact Python runtime (openpyxl, python-docx) and xlrd available.
The original files under benchmarks/literature/raw are never modified.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl
from docx import Document

ROOT = Path(__file__).resolve().parents[1] / "benchmarks/literature"
RAW = ROOT / "raw"
ROWS: list[dict] = []
REJECTED: list[dict] = []
SOURCE_URLS = {
    "chylinski2014": "https://doi.org/10.1093/nar/gku241",
    "gasiunas2020": "https://doi.org/10.1038/s41467-020-19344-1",
    "faure2018_reference": "https://doi.org/10.1080/15476286.2018.1493331",
    "faure2018_barrangou": "https://doi.org/10.1080/15476286.2018.1493331",
    "faure2018_predicted": "https://doi.org/10.1080/15476286.2018.1493331",
    "alexander2023": "https://pmc.ncbi.nlm.nih.gov/articles/PMC10278012/",
    "ziemann2025": "https://doi.org/10.1093/femsml/uqaf020",
    "dooley2021_predicted": "https://pmc.ncbi.nlm.nih.gov/articles/PMC8336202/",
    "atlas2025_predicted": "https://doi.org/10.1038/s41586-025-09298-z",
    "cas12f2026_components": "https://doi.org/10.1038/s41594-026-01788-6",
    "rfam_RF02348_seed": "https://rfam.org/family/RF02348",
    "chyou2019_predicted": "https://doi.org/10.1080/15476286.2018.1498281",
    "briner2014": "https://doi.org/10.1016/j.molcel.2014.09.019",
}


def normalize_subtype(value: str) -> str:
    value = str(value).removeprefix("CAS-")
    return {
        "IIA": "II-A",
        "IIB": "II-B",
        "IIC": "II-C",
        "VB": "V-B",
        "VE": "V-E",
        "None": "unknown",
        "UK": "unknown",
        "Unidentified": "unknown",
    }.get(value, value)


def normalize(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper().replace("U", "T")


def add(
    source: str,
    label: object,
    sequence: object,
    evidence: str,
    path: Path,
    locator: str,
    organism: object = "",
    subtype: str = "II",
    repeat: object = "",
    training: str = "unknown",
    **extra: object,
) -> None:
    sequence = normalize(sequence)
    if sequence in {"", "NA", "N/A", "NONE"} or not re.fullmatch("[ACGTRYSWKMBDHVN]+", sequence):
        REJECTED.append(
            {
                "source": source,
                "label": str(label),
                "locator": locator,
                "value": sequence,
                "reason": "missing or non-sequence text",
            }
        )
        return
    ROWS.append(
        {
            "id": f"{source}_{len(ROWS) + 1:06d}",
            "source": source,
            "label": str(label),
            "organism": str(organism or ""),
            "subtype": normalize_subtype(subtype),
            "evidence": evidence,
            "sequence": sequence,
            "length": len(sequence),
            "sequence_sha256": hashlib.sha256(sequence.encode()).hexdigest(),
            "repeat": normalize(repeat),
            "source_file": str(path.relative_to(ROOT)),
            "source_locator": locator,
            "source_url": SOURCE_URLS[source],
            "training_overlap": training,
            **extra,
        }
    )


def sheet(path: Path, name: str | None = None):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return book[name] if name else book.active


def main() -> None:
    p = next(RAW.glob("PMC4041416_supp/*Table_S4.docx"))
    for i, row in enumerate(Document(p).tables[0].rows[1:], 2):
        v = [c.text.strip() for c in row.cells]
        add(
            "chylinski2014",
            v[1],
            v[2],
            "published_mixed_prediction_and_validation",
            p,
            f"Table S4 row {i}",
            v[0],
            repeat=v[3],
            reported_length=v[5],
            training="possible_overlap_with_Briner2014",
        )

    p = next(RAW.glob("PMC7606464_supp/*MOESM5*.xlsx"))
    for i, row in enumerate(sheet(p).values, 1):
        if i == 1 or not isinstance(row[0], (int, float)) or not row[7]:
            continue
        add(
            "gasiunas2020",
            row[1],
            row[7],
            "biochemically_supported_system",
            p,
            f"Cas9_database!H{i}",
            row[3],
            repeat=row[6],
            training="source_collection_used_to_build_bundled_Type_II_models",
            clade=row[2],
            genome_label="REF_" + ("Sma" if row[1] == "Sma2" else row[1]),
        )

    # This compilation explicitly mixes old predictions and newly reported RNAs.
    # Its broad 'experimentally validated' caption is not assigned to every row.
    import xlrd

    p = next(RAW.glob("PMC6546382_supp/**/SupplementaryData1.xls"))
    book = xlrd.open_workbook(p)
    s = book.sheet_by_name("DB")
    for i in range(1, s.nrows):
        v = s.row_values(i)
        add(
            "faure2018_reference",
            f"DB_{i + 1}",
            v[1],
            "published_reference_compilation",
            p,
            f"DB row {i + 1}",
            v[3],
            subtype=str(v[0]),
            original_reference=v[2],
            training="possible_overlap_with_Briner2014",
        )
    s = book.sheet_by_name("Barrangou")
    for i in range(1, s.nrows):
        v = s.row_values(i)
        add(
            "faure2018_barrangou",
            f"Barrangou_{i + 1}",
            v[4],
            "published_reference_compilation",
            p,
            f"Barrangou row {i + 1}",
            " ".join(map(str, v[:3])),
            repeat=v[3],
            training="possible_overlap_with_Briner2014",
        )

    p = next(RAW.glob("PMC6546382_supp/**/SupplementaryData3.data"))
    for i, line in enumerate(p.read_text().splitlines(), 1):
        v = line.split()
        if len(v) != 8:
            raise ValueError(f"Unexpected Faure row {i}: {len(v)} fields")
        add(
            "faure2018_predicted",
            f"row{i}",
            v[5],
            "computational_prediction",
            p,
            f"line {i}",
            v[1],
            subtype=v[4],
            repeat=v[7],
            accession=v[0],
            reported_coordinates=v[2],
            reported_strand=v[3],
        )

    p = next(RAW.glob("PMC10278012_supp/**/Alexander*Tables*.xlsx"))
    for i, v in enumerate(sheet(p, "Table 2").values, 1):
        if i > 1 and v[5]:
            add(
                "alexander2023",
                v[0],
                v[5],
                "active_system_with_engineered_sgRNA",
                p,
                f"Table 2!F{i}",
                "metagenomic " + str(v[0]),
                training="postdates_2022_models_homology_not_excluded",
                note=(
                    "Published tracr sequence; activity of corresponding engineered sgRNA "
                    "does not establish native transcript boundaries"
                ),
            )

    p = next(RAW.glob("PMC12416283_supp/**/Dataset_3*.xlsx"))
    zi_sheet = sheet(p)
    for i, v in enumerate(zi_sheet.values, 1):
        if i > 1 and v[8]:
            add(
                "ziemann2025",
                v[0],
                v[8],
                "computational_prediction_with_selected_expression_validation",
                p,
                f"{zi_sheet.title}!I{i}",
                v[40],
                subtype="V-K_" + str(v[42]),
                repeat=v[7],
                training="source_collection_used_to_build_bundled_Type_V_models",
                reported_coordinates=f"{v[15]}-{v[16]}",
                accession=str(v[5]),
            )

    p = RAW / "AllCas9_TracrRNA_Systems.xlsx"
    for i, v in enumerate(sheet(p).values, 1):
        if i > 1 and v[18]:
            add(
                "dooley2021_predicted",
                v[0],
                v[18],
                "computational_prediction",
                p,
                f"Sheet1!S{i}",
                "",
                repeat=v[12],
                accession=v[2],
                reported_coordinates=f"{v[28]}-{v[29]}",
                reported_strand=v[20],
                training="possible_overlap_with_Gasiunas2020" if v[1] == "Corteva" else "unknown",
            )

    p = RAW / "atlas_tracr.jsonl"
    if p.exists():
        for i, line in enumerate(p.open(), 1):
            v = json.loads(line)
            add(
                "atlas2025_predicted",
                v["operon_id"],
                v["tracr"]["tracr"],
                "computational_prediction",
                p,
                f"line {i}",
                v["metadata"].get("sample_name") or v["metadata"].get("taxonomy"),
                subtype=v["summary"]["subtype"],
                repeat="",
                reported_repeats=v["repeats"],
                training="postdates_2022_models_homology_not_excluded",
                source_confidence=v["tracr"].get("confidence"),
                source_model=v["tracr"].get("cm_id"),
                license="CC BY-NC 4.0",
            )

    # Only cells with a separately labeled tetraloop/crRNA column are isolated
    # tracr components. Later rows in this same column contain entire sgRNAs,
    # a Cas12a crRNA, and target-bearing cryo-EM constructs: exclude those.
    p = next(RAW.glob("PMC13186708_supp/*MOESM4*.xlsx"))
    for i, v in enumerate(sheet(p, "b. sgRNA_sequence").values, 1):
        if i > 1 and v[0] and v[2] and v[3]:
            add(
                "cas12f2026_components",
                v[0],
                v[2],
                "published_tracr_component_of_engineered_sgRNA",
                p,
                f"b. sgRNA_sequence!C{i}",
                subtype="V-F",
                training="postdates_2022_models_homology_not_excluded",
                note=str(v[4]),
            )
    p = RAW / "RF02348_seed.fasta"
    for block in p.read_text().split(">")[1:]:
        label, sequence = block.split("\n", 1)
        add(
            "rfam_RF02348_seed",
            label,
            sequence,
            "curated_family_seed_region",
            p,
            label,
            subtype="II",
            training="likely_overlap_with_early_Type_II_references",
            note="Rfam aligned family region; not necessarily full native transcript",
        )

    p = RAW / "alternate_retrieval/S4 tracrRNA all prediction.txt"
    for i, v in enumerate(csv.DictReader(p.open()), 2):
        add(
            "chyou2019_predicted",
            v["HSP_EXTRACT_ID"],
            v["TRACR_RNA_SEQ"],
            "computational_prediction",
            p,
            f"CSV row {i}: TRACR_RNA_SEQ",
            v["ORGANISM"],
            subtype="II",
            repeat=v["REPEAT_SEQ"],
            accession=v["CONTIG_ID"],
            reported_coordinates=f"{v['TRACR_START']}-{v['TRACR_END']}",
            reported_strand="-" if v["TRACR_ORI"] == "1" else "+",
            source_subset="conservative_single_prediction"
            if v["N_PRED"] == "1"
            else "multiple_prediction_alternative",
            source_prediction_count=int(v["N_PRED"]),
            previously_reported_reference=v["CHY"] == "1",
            source_probability=float(v["PR_POS"]),
            source_tetra_t_end=v["TETRA_T_END"],
            training="possible_overlap_with_early_Type_II_references",
            license="CC BY 4.0",
            dataset_url="https://doi.org/10.6084/m9.figshare.6807203.v3",
            note=(
                "521 predicted alternatives; N_PRED=1 identifies the paper's conservative 275. "
                "CHY flags previously reported reference systems, not validated new boundaries."
            ),
        )

    extraction = RAW / "alternate_retrieval/briner2014_figure_sequences.json"
    p = RAW / "alternate_retrieval/briner_mmc1.pdf"
    for v in json.loads(extraction.read_text()):
        add(
            "briner2014",
            v["label"],
            v["sequence"],
            "published_reference_compilation",
            p,
            v["source_locator"],
            v["label"],
            subtype="II-A",
            training="source_collection_used_to_build_bundled_Type_II_models",
            extraction_file=str(extraction.relative_to(ROOT)),
            aligned_rna=v["aligned_rna"],
            extraction_note=v["normalization"],
            note=(
                "Comparative reference collection with predicted boundaries. Functional study "
                "does not validate each native transcript boundary. Figure alignment gaps removed."
            ),
        )

    ROOT.mkdir(exist_ok=True)
    with (ROOT / "catalog.jsonl").open("w") as out:
        for row in ROWS:
            out.write(json.dumps(row, ensure_ascii=False) + "\n")
    with (ROOT / "catalog.tsv").open("w", newline="") as out:
        fields = [
            "id",
            "source",
            "label",
            "organism",
            "subtype",
            "evidence",
            "length",
            "sequence",
            "repeat",
            "training_overlap",
            "source_file",
            "source_locator",
            "source_url",
            "sequence_sha256",
            "source_subset",
            "source_prediction_count",
            "previously_reported_reference",
        ]
        writer = csv.DictWriter(out, fields, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(ROWS)
    reference_rows = [r for r in ROWS if r["evidence"] != "computational_prediction"]
    with (ROOT / "reference_sequences.tsv").open("w", newline="") as out:
        writer = csv.DictWriter(out, fields, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(reference_rows)
    with (ROOT / "reference_sequences.fasta").open("w") as out:
        for row in reference_rows:
            out.write(
                f">{row['id']} {row['label']} evidence={row['evidence']}\n{row['sequence']}\n"
            )
    unique = {}
    for row in ROWS:
        unique.setdefault(row["sequence"], row)
    with (ROOT / "unique_sequences.fasta").open("w") as out:
        for seq, row in unique.items():
            out.write(f">{row['id']} sha256={row['sequence_sha256']}\n{seq}\n")
    summary = {
        "records": len(ROWS),
        "unique_sequences": len(unique),
        "by_source": dict(Counter(r["source"] for r in ROWS)),
        "by_evidence": dict(Counter(r["evidence"] for r in ROWS)),
        "rejected_rows": len(REJECTED),
        "normalization": "uppercase, remove whitespace, U to T only",
        "claim": (
            "All extractable sequences in the named retrieved tables, "
            "not all tracrRNAs ever published"
        ),
    }
    (ROOT / "catalog_summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    (ROOT / "excluded_rows.json").write_text(json.dumps(REJECTED, indent=2) + "\n")
    inputs = {r["source_file"] for r in ROWS}
    inputs.update(r["extraction_file"] for r in ROWS if r.get("extraction_file"))
    (ROOT / "input_checksums.json").write_text(
        json.dumps(
            {p: hashlib.sha256((ROOT / p).read_bytes()).hexdigest() for p in sorted(inputs)},
            indent=2,
        )
        + "\n"
    )
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
